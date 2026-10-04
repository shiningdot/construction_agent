#!/usr/bin/env python3
"""V2 四格式构建：EN DOCX + EN PDF + ZH DOCX + ZH PDF
- DOCX: python-docx（中文设宋体东亚字体）
- PDF: LaTeX + tectonic（CJK 用 Noto Sans SC）
- 图片占位符 [图X占位...] 渲染为可见占位框"""
import re, os, subprocess
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = '/home/suntianhao/construction_agent/.claude/worktrees/classifier-mod/docs/paper_v2'
TECTONIC = '/home/suntianhao/TEM/tectonic'
FIGDIR = '/home/suntianhao/construction_agent/.claude/worktrees/classifier-mod/docs/paper_v2/figures'
os.makedirs(FIGDIR, exist_ok=True)

SRC = {
    'en': f'{BASE}/eswa_paper_v2.md',
    'zh': f'{BASE}/eswa_paper_v2_zh.md',
}

# 正文插图：![](figures/figN_xxx.png)，相对 BASE。旧的 [图N占位] 分支保留作兜底。
IMG_RE = re.compile(r'^!\[[^\]]*\]\(([^)]+)\)\s*$')

def img_abspath(rel):
    return os.path.normpath(os.path.join(BASE, rel))

def img_width_inches(path, max_w=6.0, max_h=7.4):
    """按纵横比定宽，避免 fig1 这类竖长图（9.6×12）撑爆一页。"""
    try:
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
    except Exception:
        return max_w
    ratio = h / max(1, w)
    return max_h / ratio if max_w * ratio > max_h else max_w

def clean(text):
    text = re.sub(r'<!--[\s\S]*?-->', '', text)
    text = re.sub(r'^> .*\n', '', text, flags=re.M)
    return text

def cjk_runs(para, s, font_zh='宋体'):
    """按中文片段切分并设置东亚字体"""
    parts = re.split(r'([\u4e00-\u9fff\u3000-\u303f\uff00-\uffef]+)', s)
    for part in parts:
        if not part:
            continue
        r = para.add_run(part)
        if re.match(r'[\u4e00-\u9fff]', part):
            r.font.name = font_zh
            r._element.rPr.rFonts.set(qn('w:eastAsia'), font_zh)

# ══════════════ DOCX ══════════════
def build_docx(src, out, is_zh):
    text = clean(open(src, encoding='utf-8').read())
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(10.5)
    if is_zh:
        st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    lines = text.split('\n')
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s == '' or s == '---':
            i += 1
            continue
        if s.startswith('|') and i + 1 < len(lines) and lines[i+1].strip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            header, body = rows[0], rows[2:]
            table = doc.add_table(rows=1+len(body), cols=len(header))
            table.style = 'Light Grid Accent 1'
            for j, hcell in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.text = hcell.replace('**', '')
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for ri, r in enumerate(body, start=1):
                r = r + [''] * (len(header) - len(r))
                for j, c in enumerate(r[:len(header)]):
                    table.rows[ri].cells[j].text = c.replace('**', '')
            continue
        if s.startswith('# '):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(s[2:].strip()); r.bold = True; r.font.size = Pt(17)
            i += 1
            continue
        if s.startswith('## '):
            p = doc.add_heading(level=1)
            _add_runs(p, s[3:].strip(), is_zh)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.size = Pt(14)
            i += 1
            continue
        if s.startswith('### '):
            p = doc.add_heading(level=2)
            _add_runs(p, s[4:].strip(), is_zh)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue
        if s.startswith('#### '):  # 四级标题（3.2.1 一类），缺这一支会被整行丢掉
            p = doc.add_heading(level=3)
            _add_runs(p, s[5:].strip(), is_zh)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue
        m = IMG_RE.match(s)
        if m:
            path = img_abspath(m.group(1))
            if os.path.exists(path):
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(path, width=Inches(img_width_inches(path)))
            i += 1
            continue
        if re.match(r'^\[图\d+占位', s):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            r = p.add_run('▢ ' + s)
            r.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
            r.italic = True
            i += 1
            continue
        if s.startswith('**Figure') or s.startswith('**图'):
            p = doc.add_paragraph()
            _add_runs(p, s, is_zh)
            for run in p.runs:
                run.bold = True
            i += 1
            continue
        if s.startswith('- '):
            while i < len(lines) and lines[i].strip().startswith('- '):
                p = doc.add_paragraph(style='List Bullet')
                _add_runs(p, lines[i].strip()[2:], is_zh)
                i += 1
            continue
        if re.match(r'^[A-Z][\wÀ-ſ., \-]+\(\d{4}', s) and 'References' in ''.join(lines[max(0,i-3):i]):
            pass  # 参考文献条目按普通段落
        p = doc.add_paragraph()
        _add_runs(p, s, is_zh)
        i += 1
    doc.save(out)
    print(f'DOCX: {out}')

def _add_runs(p, s, is_zh):
    tokens = re.split(r'(\*\*.+?\*\*|\*[^*]+?\*|\$[^$]+\$)', s)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            if is_zh:
                r = p.add_run(); r.bold = True
                cjk_sub(r, tok[2:-2], '宋体')
            else:
                r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = p.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith('$') and tok.endswith('$'):
            _add_math(p, tok[1:-1])
        else:
            if is_zh:
                cjk_runs(p, tok)
            else:
                p.add_run(tok)


# DOCX 路径不经过 esc()，此前 $T = \{t_1, \dots, t_N\}$ 是原样输出的——读者在 Word
# 里看到的是美元号和反斜杠命令本身。Word 没有轻量的行内公式写入口（OMML 要手工拼
# XML），退而用「斜体正文 + 真下标」，足以表达 t_i、x_ij 这类记号。
_MATH_MAP = {'\\dots': '…', '\\in': '∈', '\\iota': 'ι', '\\{': '{', '\\}': '}'}


def _add_math(p, tex):
    s = tex
    for k, v in _MATH_MAP.items():
        s = s.replace(k, v)
    s = re.sub(r'\\text\{([^}]*)\}', r'\1', s)   # \text{explicit} -> explicit
    # 剩下的 \ 是没映射到的命令，去掉反斜杠总比原样露出好（现只有这几条，实际不会走到）
    s = s.replace('\\', '')
    # _ 之后跟 {...} 取整组作下标，否则取一个字符
    for part in re.split(r'(_(?:\{[^}]*\}|.))', s):
        if not part:
            continue
        if part.startswith('_'):
            r = p.add_run(part[1:].strip('{}'))
            r.italic = True
            r.font.subscript = True
        else:
            r = p.add_run(part)
            r.italic = True

def cjk_sub(r, s, font_zh):
    r.text = s
    r.font.name = font_zh
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font_zh)

# ══════════════ LaTeX ══════════════
# 标题开头的手写编号（"1. "、"2.1 "、"3.2.1 "）。markdown 里手写、LaTeX 又自动编号，
# 不剥掉就渲染成 "2 1. Introduction"。
NUMPFX = re.compile(r'^\d+(?:\.\d+)*\.?\s+')


def esc(s):
    # markdown 里已经写好的 $...$ 数学片段必须整段保护：下面的 replace 会把 _ 转义成
    # \_，于是 $t_i$、$x_{ij}$ 变成 $t\_i$，渲染成字面下划线而不是下标（实测 §3.1 全段
    # 如此）。先抽成占位符，所有转义做完再原样放回。全文 $ 配对已核（两版各 10 处，
    # 无奇数行），不会误吞正文。
    math = []

    def _stash(m):
        math.append(m.group(0))
        return f'\x00{len(math) - 1}\x00'

    s = re.sub(r'\$[^$\n]+\$', _stash, s)
    # 反引号代码片段（文件路径、字段名）是正文里最长的不可断串，实测把行推出页外
    # 41pt；先原样插断点，之后再轮到 _ % # 这些转义
    s = re.sub(r'`([^`]+)`',
               lambda m: '`' + re.sub(r'([/._-])', r'\1\\allowbreak{}', m.group(1)) + '`', s)
    s = s.replace('%', r'\%').replace('&', r'\&').replace('#', r'\#').replace('_', r'\_')
    s = s.replace('κ', r'$\kappa$').replace('Δ', r'$\Delta$')
    s = s.replace('→', r'$\rightarrow$').replace('×', r'$\times$')
    s = s.replace('≈', r'$\approx$').replace('–', '--').replace('—', '---')
    s = s.replace('’', "'").replace('“', "``").replace('”', "''")
    # 圈码 ①-⑫（U+2460 起）在 Latin Modern 与 \textit 下都没有字形，实测两版各 12 个
    # 全部渲染成豆腐块（"Missing character" 警告）。正文里它们只是列表序号，换成
    # (1)-(12) 既规避字体依赖，也是学术论文的通行写法。
    s = re.sub(r'[\u2460-\u2473]', lambda m: f'({ord(m.group(0)) - 0x245F})', s)
    s = re.sub(r'\*\*(.+?)\*\*', r'\\textbf{\1}', s)
    s = re.sub(r'(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)', r'\\textit{\1}', s)
    # 中文串在 TeX 里是「无空格的长单词」，不插断点则整串不可断——正文的普通段落
    # 尚可溢出到页外，p{} 列里就是必然溢出。逐字插 \allowbreak（不产生空格的合法
    # 断点）。不用 \XeTeXlinebreaklocale：它需要 ICU 断行数据，本机 tectonic 缺，
    # 会直接以 "failed to create linebreak iterator, status=2" 中止编译。
    s = re.sub(r'([一-鿿　-〿＀-￯]+)',
               lambda m: r'{\cjkfont ' + r'\allowbreak '.join(m.group(1)) + '}', s)
    # 参考文献里的 DOI 没有空格可断，实测最长被推出页面 156pt；在 / . - 之后补断点
    # （\allowbreak{} 后面跟空组，免得 TeX 把紧跟的空格一并吃掉）
    s = re.sub(r'https?://[^\s]+',
               lambda m: re.sub(r'([/.-])', r'\1\\allowbreak{}', m.group(0)), s)
    return re.sub(r'\x00(\d+)\x00', lambda m: math[int(m.group(1))], s)

def _disp_pt(s):
    """估算字符串宽度（pt）：11pt 字号下 CJK 约 1em(11pt)、ASCII 约 0.5em(5.5pt)。"""
    return sum(11.0 if ord(c) > 0x2000 else 5.5 for c in s)


def _unbreakable(s):
    """切出「不可断单元」，用来算列宽的硬下限。

    不能直接 split()：中文单元格里没有空格，整段定义会被当成一个上百 pt 的
    「单词」，五列的下限之和必然超过页宽，于是永远走等比压缩分支、每列都被压到
    自己的下限以下，最长单词逐行溢出（中文表 2 实测每行 2–5pt，且对 TEXTW 改
    动完全不敏感）。esc() 给中文逐字插了 \\allowbreak，所以中文的可断单元就是
    单字；ASCII 段仍按整词（词内无可断点）。
    """
    toks, cur = [], ''
    for ch in s:
        if ord(ch) > 0x2000:  # CJK 及全角标点：逐字可断
            if cur:
                toks.append(cur)
                cur = ''
            toks.append(ch)
        elif ch.isspace():
            if cur:
                toks.append(cur)
                cur = ''
        else:
            cur += ch
    if cur:
        toks.append(cur)
    return toks


def _cell_pt(c):
    """单元格宽度，含 markdown 粗体的单元格整体加权。

    esc() 把 **x** 变成 \\textbf{x}，实测粗体比常规字宽约 12%。表 10 末行七列
    全粗体，(operational) 这类 13 字符的词按常规宽度估就差 5.58pt，逐行溢出。
    """
    return _disp_pt(c) * (1.12 if '**' in c else 1.0)


def colspec(rows, ncol):
    """分配列宽，返回 LaTeX 列规格串。

    两步走：每列先取「最长不可断单元放得下」作硬下限，剩余宽度再按内容长度加权分。
    只按内容加权不够——表 3 的 Categories 列数据全是两位数，加权后只剩 6.5%
    宽（29pt），而表头 "Categories" 要 45pt，那一行就固定溢出 15pt。
    长度在 esc() 之前量，否则每个中文字符都会被 {\\cjkfont ...} 包装撑长。
    """
    TEXTW = 455.0  # A4 减 2×2.4cm 页边实测 \textwidth≈459pt，留 4pt 吸收估算误差
    data = [r for k, r in enumerate(rows) if k != 1]  # k=1 是 |---|---| 分隔行
    avail = TEXTW - (2 * 2 * ncol + 0.4 * (ncol + 1))  # 扣掉 \tabcolsep(2pt) 与竖线

    mins, wants = [], []
    for j in range(ncol):
        # 空单元格是「沿用上一行」的续行标记，不计入——否则表 2 的层名列（24 行
        # 里只有 5 行有值）会被压到下限，"M — Management foundation" 断成 5 行。
        cells = [r[j] for r in data if j < len(r) and r[j].strip()] or ['']
        longest = max((_disp_pt(t) * (1.12 if '**' in c else 1.0)
                       for c in cells for t in _unbreakable(c)), default=1.0)
        mins.append(longest + 4)  # 留 4pt 内边距
        w = (0.5 * max(map(_cell_pt, cells))
             + 0.5 * sum(map(_cell_pt, cells)) / len(cells))
        if len(cells) < 0.4 * len(data):
            w *= 1.5  # 稀疏的分组列装的是整体标签，宽一点才不会被断字
        wants.append(max(1.0, w))

    if sum(mins) >= avail:  # 列太多或单词太长，硬下限已超页宽，只好等比压缩
        widths = [m * avail / sum(mins) for m in mins]
    else:
        spare, tw = avail - sum(mins), sum(wants)
        widths = [m + spare * w / tw for m, w in zip(mins, wants)]
    return ''.join(r'>{\raggedright\arraybackslash}p{'
                   + f'{w / TEXTW:.4f}' + r'\textwidth}' for w in widths)


def build_latex(src, out_tex, is_zh):
    text = clean(open(src, encoding='utf-8').read())
    lines = text.split('\n')
    out = [r"""\documentclass[11pt]{article}
% 必须显式写 a4paper：geometry 不指定纸张时沿用 article 的 letterpaper 默认，
% 成品实测是 612x792pt（US Letter），\textheight 只有 656pt——既是排版警告的来源，
% 也是投稿格式错误。A4 下 \textwidth≈459pt、\textheight≈706pt。
\usepackage[a4paper,margin=2.4cm]{geometry}
\usepackage{graphicx}
\usepackage{amsmath,amssymb}
\usepackage{booktabs}
\usepackage[hidelinks]{hyperref}
\usepackage{microtype}
\usepackage{float}
\usepackage{longtable}
\usepackage{array}
\usepackage{fontspec}
\newfontfamily\cjkfont{Noto Sans SC}[Scale=0.92]
% 中文段落只有 \allowbreak 断点、没有可伸缩空白，默认 tolerance 下 TeX 找不到
% 满意解就只能让整行溢出（实测最大 93pt）。放宽 tolerance 并允许 emergency
% 伸缩——它们只在原本无解的位置生效，不会把已经断好的行改差。
\tolerance=3000
\emergencystretch=1.5em
% longtable（表 2 有 25 行）与图注、节标题交替出现时，TeX 在 \output 阶段会尝试
% 组装一个高于 \textheight 的页面盒并报 "vbox too high"，实测中文版 232pt。加了
% \raggedbottom 后底部不强行对齐，该警告消失；实际渲染本就没有内容越界（18 页逐页
% 量过正文块底边），所以这是消噪而非修错。
\raggedbottom
""" + (r'\title{' + esc('CAIEMS：面向施工事故叙述结构化处理的探索性 LLM '
                     '数据工程管道') + '}' if is_zh
       else r'\title{CAIEMS: An Exploratory LLM-Driven Pipeline for Structuring Construction Accident Narratives into a Queryable Database}')
     + r"""
\author{Tianhao Sun\textsuperscript{1}, [Supervisor name]\textsuperscript{1,*}}
\date{2026-09-14}
\begin{document}
\maketitle
"""]
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s == '' or s == '---':
            i += 1
            continue
        if s.startswith('|') and i + 1 < len(lines) and lines[i+1].strip().startswith('|'):
            raw = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                raw.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            rows = [[esc(c) for c in r] for r in raw]
            header = rows[0]
            ncol = len(header)
            body = [r + [''] * (ncol - len(r)) for r in rows[2:]]
            # longtable 而非 table+tabular：表 2 有 25 行，浮动体装不下就被截断；
            # 列宽用 p{} 受控，否则近百字符的定义列会把页面撑爆。
            # 列间留白收到 2pt：表 10 有 7 列，3pt 时 \tabcolsep 就吃掉 45pt，
            # 内容区只剩 410pt，末行七列全粗体差 4.28pt 放不下。2pt 让出 14pt。
            out.append(r'\begingroup\setlength{\tabcolsep}{2pt}\small')
            out.append(r'\begin{longtable}{' + colspec(raw, ncol) + '}')
            out.append(r'\hline')
            out.append(' & '.join(header) + r' \\')
            out.append(r'\hline\endfirsthead')
            out.append(r'\hline')
            out.append(' & '.join(header) + r' \\')
            out.append(r'\hline\endhead')
            for r in body:
                out.append(' & '.join(r) + r' \\')
            out.append(r'\hline')
            out.append(r'\end{longtable}\endgroup')
            continue
        if s.startswith('# '):
            i += 1
            continue
        if s.startswith('## '):
            # markdown 标题自带手写编号（"1. Introduction"、"2.1 ..."），而 \section 由
            # LaTeX 自动编号，两者会叠成 "2 1. Introduction"。剥掉手写号交给自动编号，
            # 结果与手写号一致（摘要/参考文献/附录例外，见下）。
            sec = NUMPFX.sub('', s[3:].strip())
            if sec.startswith('References') or sec.startswith('参考文献'):
                out.append(r'\section*{' + esc(sec) + '}' + '\n' + r'\begin{enumerate}')
                i += 1
                while i < len(lines) and lines[i].strip() == '':
                    i += 1
                while i < len(lines) and lines[i].strip() != '':
                    ref = lines[i].strip()
                    if not re.match(r'^(\*V2|V2)', ref):
                        out.append(r'\item ' + esc(ref))
                    i += 1
                out.append(r'\end{enumerate}')
                continue
            if sec.startswith('Appendix') or sec.startswith('附录'):
                # \appendix 之后自动编号就变成 A、B…，标题里再留 "Appendix A. " 会
                # 叠成 "A Appendix A. Prompt Templates"，所以这一支去掉字母号并用星号版。
                out.append(r'\appendix')
                sec = re.sub(r'^(Appendix|附录)\s+[A-Z\d]+\.?\s*', '', sec)
                out.append(r'\section*{' + esc(sec) + '}')
                i += 1
                continue
            star = '*' if sec in ('Abstract', '摘要') else ''
            out.append(r'\section' + star + '{' + esc(sec) + '}')
            i += 1
            continue
        if s.startswith('### '):
            out.append(r'\subsection{' + esc(NUMPFX.sub('', s[4:].strip())) + '}')
            i += 1
            continue
        if s.startswith('#### '):
            out.append(r'\subsubsection{'
                       + esc(NUMPFX.sub('', s[5:].strip())) + '}')
            i += 1
            continue
        m = IMG_RE.match(s)
        if m:
            path = img_abspath(m.group(1))
            if os.path.exists(path):
                # 不用 h(ere)：图 1 是 1.24 的竖版，0.92\textwidth 下实高约 522pt，
                # 就地插入会和前后段落凑成一个高于 \textheight 的盒子，正是 "vbox
                # too high" 的成因。t/b/p 让 TeX 自己找位置，宽度已由 \textwidth 控死。
                out.append(r'\begin{figure}[tbp]' + '\n' + r'\centering' + '\n'
                           + r'\includegraphics[width=0.92\textwidth,'
                           + r'height=0.80\textheight,keepaspectratio]{' + path + r'}'
                           + '\n' + r'\end{figure}')
            i += 1
            continue
        if re.match(r'^\[图\d+占位', s):
            out.append(r'\begin{center}\fbox{\parbox{0.8\textwidth}{\color{gray}' + esc(s) + r'}}\end{center}')
            out.append(r'\usepackage{xcolor}' if False else '')
            i += 1
            continue
        if s.startswith('**Figure') or s.startswith('**图'):
            out.append(r'\textbf{' + esc(s.strip('*')) + '}')
            i += 1
            continue
        if s.startswith('*Notes.*') or s.startswith('*注。*'):
            out.append(r'\textit{' + esc(s) + r'}\\')
            i += 1
            continue
        if s.startswith('- '):
            out.append(r'\begin{itemize}')
            while i < len(lines) and lines[i].strip().startswith('- '):
                out.append(r'\item ' + esc(lines[i].strip()[2:]))
                i += 1
            out.append(r'\end{itemize}')
            continue
        para = []
        while i < len(lines) and lines[i].strip() not in ('', '---') \
                and not lines[i].strip().startswith(('#', '|', '**Figure', '**图', '- ')) \
                and not re.match(r'^\[图\d+占位', lines[i].strip()) \
                and not IMG_RE.match(lines[i].strip()):
            para.append(lines[i])
            i += 1
        if para:
            out.append(esc(' '.join(para).strip()))
            out.append('')
        else:
            i += 1
    out.append(r'\end{document}')
    tex = '\n'.join(out)
    # 补充 xcolor
    tex = tex.replace(r'\usepackage{fontspec}', r'\usepackage{fontspec}' + '\n' + r'\usepackage{xcolor}')
    open(out_tex, 'w', encoding='utf-8').write(tex)
    print(f'TeX: {out_tex}')
    r = subprocess.run([TECTONIC, out_tex], capture_output=True, text=True,
                       cwd='/home/suntianhao/construction_agent/.claude/worktrees/classifier-mod')
    pdf = out_tex.replace('.tex', '.pdf')
    ok = os.path.exists(pdf)
    if not ok:
        for ln in r.stderr.split('\n')[-8:]:
            if 'error' in ln.lower():
                print(ln[:200])
    print(f'PDF: {pdf} {"✓" if ok else "✗ 失败"}')
    return ok

if __name__ == '__main__':
    build_docx(SRC['en'], f'{BASE}/eswa_paper_v2_en.docx', False)
    build_docx(SRC['zh'], f'{BASE}/eswa_paper_v2_zh.docx', True)
    build_latex(SRC['en'], f'{BASE}/eswa_paper_v2_en.tex', False)
    build_latex(SRC['zh'], f'{BASE}/eswa_paper_v2_zh.tex', True)
    print('全部完成')
