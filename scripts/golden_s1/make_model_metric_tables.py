"""生成两张论文三线表。

表 1：四个模型 × 六种策略的平均 F1（Ave / Ave+）。
表 2：最优模型 DeepSeek V4 Flash 的逐类 F1、召回、精度、准确性明细。

数据源：experiments/golden_standard/figures/figure_data.json（专家 GT 口径）。
输出：DOCX（主）、Markdown、CSV，放在图目录下。

用法：
    python scripts/golden_s1/make_model_metric_tables.py
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from make_definition_table_zh import (  # noqa: E402
    _set_cell_bottom_border,
    _set_run_font,
    _set_table_fixed_width,
    _set_table_three_line_borders,
)

MODEL_ORDER = ["glm5.2", "flash", "hy3", "pro"]
STRATEGY_ORDER = ["S0", "S1", "S2", "S3", "S4", "S5"]
STRATEGY_ZH = {
    "S0": "S0 零样本",
    "S1": "S1 少样本",
    "S2": "S2 逐层推理",
    "S3": "S3 计划与求解",
    "S4": "S4 CAICS-CoT",
    "S5": "S5 CAICS-CoT+",
}
METRICS = [
    ("F1", "F1 分数"),
    ("R", "召回"),
    ("P", "精度"),
    ("Acc", "准确性"),
]
NOTE_AVG = (
    "注：Ave 为 19 类未加权宏平均；Ave+ 为 GT 正例数 ≥5 的 18 类（排除 C2）。"
    "真值为专家 GT（N=300，平均 7.85 编码/条）。"
    "每列最优值加粗，次优值加下划线。"
)
NOTE_FLASH = (
    "注：Ave / Ave+ 定义同上；数值为百分比。"
    "每个指标块内各列最优值加粗、次优值加下划线。"
    "DeepSeek V4 Flash 为四模型平均 Ave+ F1 最高者（81.71%）。"
)


def _set_section(doc: Document, landscape: bool) -> None:
    """A4 纵向或横向，留出表格宽度。"""
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = section.right_margin = Cm(1.5)
        section.top_margin = section.bottom_margin = Cm(1.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(2.0)


def _add_title(doc: Document, text: str, size: float = 10.5) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(3)
    _set_run_font(paragraph.add_run(text), size=size, bold=True)


def _add_note(doc: Document, text: str, size: float = 9) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    _set_run_font(paragraph.add_run(text), size=size)


def _style_cell(
    cell,
    text: str,
    size: float,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = text
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    for run in paragraph.runs:
        _set_run_font(run, size=size)


def _repeat_header(table) -> None:
    row_properties = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    row_properties.append(header_flag)


def _set_group_bottom_border(table, row_index: int, sz: int = 4) -> None:
    """组间分隔线：0.5 pt。"""
    for cell in table.rows[row_index].cells:
        _set_cell_bottom_border(cell, sz=sz)


def _highlight_column(table, col: int, start_row: int, end_row: int) -> None:
    """块内某列：最优加粗，次优加下划线。"""
    values: list[tuple[float, int]] = []
    for row in range(start_row, end_row + 1):
        text = table.cell(row, col).text.strip()
        if not text:
            continue
        values.append((float(text), row))
    if not values:
        return
    unique = sorted({value for value, _ in values}, reverse=True)
    best = unique[0]
    second = unique[1] if len(unique) > 1 else None
    for value, row in values:
        for run in table.cell(row, col).paragraphs[0].runs:
            if value == best:
                run.bold = True
            elif second is not None and value == second:
                run.underline = True


def _merge_first_column(table, labels: list[str], group_size: int, size: float) -> None:
    for group, label in enumerate(labels):
        start = 1 + group * group_size
        end = start + group_size - 1
        merged = table.cell(start, 0).merge(table.cell(end, 0))
        merged.text = label
        merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in merged.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                _set_run_font(run, size=size)


def _build_table_a(doc_data: dict) -> tuple[list[str], list[list[str]]]:
    headers = ["模型", "策略", "Ave F1", "Ave+ F1"]
    rows: list[list[str]] = []
    for model in MODEL_ORDER:
        for strategy in STRATEGY_ORDER:
            cell = doc_data["cells"][strategy][model]
            rows.append(
                [
                    doc_data["labels"]["model_display"][model],
                    STRATEGY_ZH[strategy],
                    f"{cell['ave']['F1'] * 100:.2f}",
                    f"{cell['ave_plus']['F1'] * 100:.2f}",
                ]
            )
    return headers, rows


def write_table_a(doc_data: dict, path: Path) -> None:
    headers, rows = _build_table_a(doc_data)
    doc = Document()
    _set_section(doc, landscape=False)
    _add_title(doc, "表 1. 四个模型 × 六种策略的平均 F1 分数比较（%）")

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_fixed_width(table, 16.0)

    for index, header in enumerate(headers):
        _style_cell(
            table.rows[0].cells[index], header, 10, WD_ALIGN_PARAGRAPH.CENTER
        )
        _set_cell_bottom_border(table.rows[0].cells[index], sz=6)

    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            align = (
                WD_ALIGN_PARAGRAPH.LEFT
                if index < 2
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            _style_cell(cells[index], value, 10, align)

    _merge_first_column(
        table,
        [doc_data["labels"]["model_display"][m] for m in MODEL_ORDER],
        group_size=6,
        size=10,
    )

    # 模型组之间的 0.5 pt 分隔线
    for group in range(len(MODEL_ORDER) - 1):
        _set_group_bottom_border(table, 1 + (group + 1) * 6 - 1, sz=4)

    _set_table_three_line_borders(table)

    widths = [Cm(3.2), Cm(4.0), Cm(4.4), Cm(4.4)]
    for table_row in table.rows:
        for index, width in enumerate(widths):
            table_row.cells[index].width = width

    _highlight_column(table, 2, 1, 24)
    _highlight_column(table, 3, 1, 24)
    _repeat_header(table)
    _add_note(doc, NOTE_AVG, size=9)
    doc.save(path)


def _build_table_b(doc_data: dict) -> tuple[list[str], list[str], list[list[str]]]:
    codes = [
        code
        for layer in doc_data["labels"]["layers"]
        for code in layer["codes"]
    ]
    headers = ["指标", "策略"] + codes + ["Ave", "Ave+"]
    rows: list[list[str]] = []
    for metric_key, metric_label in METRICS:
        for strategy in STRATEGY_ORDER:
            cell = doc_data["cells"][strategy]["flash"]
            values = [metric_label, STRATEGY_ZH[strategy]]
            values.extend(
                f"{cell['per'][metric_key][code] * 100:.2f}" for code in codes
            )
            values.extend(
                [
                    f"{cell['ave'][metric_key] * 100:.2f}",
                    f"{cell['ave_plus'][metric_key] * 100:.2f}",
                ]
            )
            rows.append(values)
    return codes, headers, rows


def write_table_b(doc_data: dict, path: Path) -> None:
    codes, headers, rows = _build_table_b(doc_data)
    doc = Document()
    _set_section(doc, landscape=True)
    _add_title(
        doc,
        "表 2. DeepSeek V4 Flash 在六种策略下的逐类 F1、召回、精度与准确性（%）",
    )

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_fixed_width(table, 25.95)

    for index, header in enumerate(headers):
        _style_cell(
            table.rows[0].cells[index], header, 8, WD_ALIGN_PARAGRAPH.CENTER
        )
        _set_cell_bottom_border(table.rows[0].cells[index], sz=6)

    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            align = (
                WD_ALIGN_PARAGRAPH.LEFT
                if index < 2
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            _style_cell(cells[index], value, 8, align)

    _merge_first_column(
        table,
        [label for _, label in METRICS],
        group_size=6,
        size=8,
    )

    # 指标块之间的 0.5 pt 分隔线
    for group in range(len(METRICS) - 1):
        _set_group_bottom_border(table, 1 + (group + 1) * 6 - 1, sz=4)

    _set_table_three_line_borders(table)

    widths = [Cm(1.7), Cm(2.0)] + [Cm(1.05)] * len(codes) + [
        Cm(1.15),
        Cm(1.15),
    ]
    for table_row in table.rows:
        for index, width in enumerate(widths):
            table_row.cells[index].width = width

    for group in range(len(METRICS)):
        start = 1 + group * 6
        end = start + 5
        for col in range(2, len(headers)):
            _highlight_column(table, col, start, end)

    _repeat_header(table)
    _add_note(doc, NOTE_FLASH, size=8)
    doc.save(path)


def write_csv(headers: list[str], rows: list[list[str]], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def write_markdown(
    title: str, headers: list[str], rows: list[list[str]], note: str, path: Path
) -> None:
    lines = [
        f"**{title}**",
        "",
        "| " + " | ".join(headers) + " |",
        "|" + "|".join(["---"] * len(headers)) + "|",
    ]
    for row in rows:
        lines.append("| " + " | ".join(row) + " |")
    lines.extend(["", note, ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="生成模型/策略指标三线表")
    parser.add_argument(
        "--data-json",
        default=str(ROOT / "experiments/golden_standard/figures/figure_data.json"),
    )
    parser.add_argument(
        "--out-dir",
        default=str(ROOT / "experiments/golden_standard/figures"),
    )
    args = parser.parse_args()

    data_path = Path(args.data_json).expanduser().resolve()
    out_dir = Path(args.out_dir).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    doc_data = json.loads(data_path.read_text(encoding="utf-8"))

    headers_a, rows_a = _build_table_a(doc_data)
    codes_b, headers_b, rows_b = _build_table_b(doc_data)

    table_a_docx = out_dir / "table_model_strategy_avg_f1.docx"
    table_b_docx = out_dir / "table_flash_four_metrics.docx"
    write_table_a(doc_data, table_a_docx)
    write_table_b(doc_data, table_b_docx)

    write_csv(headers_a, rows_a, out_dir / "table_model_strategy_avg_f1.csv")
    write_markdown(
        "表 1. 四个模型 × 六种策略的平均 F1 分数比较（%）",
        headers_a,
        rows_a,
        NOTE_AVG,
        out_dir / "table_model_strategy_avg_f1.md",
    )
    write_csv(headers_b, rows_b, out_dir / "table_flash_four_metrics.csv")
    write_markdown(
        "表 2. DeepSeek V4 Flash 在六种策略下的逐类 F1、召回、精度与准确性（%）",
        headers_b,
        rows_b,
        NOTE_FLASH,
        out_dir / "table_flash_four_metrics.md",
    )

    for path in (
        table_a_docx,
        table_b_docx,
        out_dir / "table_model_strategy_avg_f1.csv",
        out_dir / "table_model_strategy_avg_f1.md",
        out_dir / "table_flash_four_metrics.csv",
        out_dir / "table_flash_four_metrics.md",
    ):
        print(f"  wrote {path} ({path.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
