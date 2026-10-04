#!/usr/bin/env python3
"""V1 → V2 彩色对比文档生成器：红=删除（删除线），绿=新增
输出：docs/paper_v2/v1_v2_对比.html（浏览器）+ v1_v2_对比.docx（Word）"""
import re, difflib, html as html_mod
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor

# 本仓库根（脚本在 scripts/ 下，上溯一级），worktree 与主仓库都成立。
# 原先四处写死 `/.claude/worktrees/classifier-mod/`。
ROOT = Path(__file__).resolve().parents[1]
V1 = str(ROOT / 'docs' / 'paper' / 'eswa_paper.md')
V2 = str(ROOT / 'docs' / 'paper_v2' / 'eswa_paper_v2.md')
OUT_HTML = str(ROOT / 'docs' / 'paper_v2' / 'v1_v2_对比.html')
OUT_DOCX = str(ROOT / 'docs' / 'paper_v2' / 'v1_v2_对比.docx')

RED = '#C0392B'
GREEN = '#1E8449'
GRAY = '#7F8C8D'


def load(path):
    text = open(path, encoding='utf-8').read()
    text = re.sub(r'<!--[\s\S]*?-->', '', text)      # 去注释
    text = re.sub(r'^> .*\n', '', text, flags=re.M)   # 去注记行
    return text


def split_sections(text):
    """按 ## 标题分节，返回 [(标题, [段落]), ...]"""
    parts = re.split(r'(^## .+$)', text, flags=re.M)
    sections = []
    for i in range(1, len(parts), 2):
        title = parts[i].strip().lstrip('#').strip()
        body = parts[i + 1] if i + 1 < len(parts) else ''
        paras = [p.strip() for p in body.split('\n\n') if p.strip()]
        sections.append((title, paras))
    return sections


v1_text, v2_text = load(V1), load(V2)
s1, s2 = split_sections(v1_text), split_sections(v2_text)
titles1 = [t for t, _ in s1]
titles2 = [t for t, _ in s2]

# 节标题配对（容忍重命名）
def match_sections(t1, t2):
    pairs, used = [], set()
    for i, a in enumerate(t1):
        best, best_j = 0, -1
        for j, b in enumerate(t2):
            if j in used:
                continue
            r = difflib.SequenceMatcher(None, a[:60], b[:60]).ratio()
            if r > best:
                best, best_j = r, j
        if best > 0.45:
            pairs.append((i, best_j, best))
            used.add(best_j)
        else:
            pairs.append((i, -1, 0))
    return pairs

pairs = match_sections(titles1, titles2)

# 构建 diff 记录: [('del', text), ('add', text), ('same', text)]
diff_sections = []
for i, j, ratio in pairs:
    title = titles1[i]
    old_paras = s1[i][1] if i < len(s1) else []
    new_paras = s2[j][1] if j >= 0 and j < len(s2) else []
    entries = []
    if j < 0:
        for p in old_paras:
            entries.append(('del', p))
    else:
        sm = difflib.SequenceMatcher(None, old_paras, new_paras)
        for tag, a1, a2, b1, b2 in sm.get_opcodes():
            if tag == 'equal':
                for p in old_paras[a1:a2]:
                    entries.append(('same', p))
            elif tag == 'delete':
                for p in old_paras[a1:a2]:
                    entries.append(('del', p))
            elif tag == 'insert':
                for p in new_paras[b1:b2]:
                    entries.append(('add', p))
            elif tag == 'replace':
                for p in old_paras[a1:a2]:
                    entries.append(('del', p))
                for p in new_paras[b1:b2]:
                    entries.append(('add', p))
    diff_sections.append((title, entries))

# ── HTML 输出 ──
h = []
h.append('<!DOCTYPE html><html><head><meta charset="utf-8"><title>V1→V2 对比</title>')
h.append('<style>body{font-family:"Times New Roman",serif;max-width:1000px;margin:30px auto;line-height:1.7}'
         'h1{color:#2C3E50} h2{color:#2C3E50;border-bottom:2px solid #2C3E50;padding-top:20px}'
         '.del{color:#C0392B;text-decoration:line-through;background:#FDEDEC;padding:1px 3px}'
         '.add{color:#1E8449;background:#EAFAF1;padding:1px 3px;font-weight:normal}'
         '.same{color:#555} .tag{font-size:12px;padding:1px 6px;border-radius:3px;margin-right:6px}'
         '.tag-del{background:#C0392B;color:#fff}.tag-add{background:#1E8449;color:#fff}</style></head><body>')
h.append('<h1>V1 → V2 修订对比（红=删除，绿=新增）</h1>')
h.append('<p style="color:#7F8C8D">V1: docs/paper/eswa_paper.md ｜ V2: docs/paper_v2/eswa_paper_v2.md ｜ 生成日期 2026-09-14</p>')
stats = {'del': 0, 'add': 0, 'same': 0}
for title, entries in diff_sections:
    h.append(f'<h2>{html_mod.escape(title)}</h2>')
    for tag, text in entries:
        stats[tag] = stats.get(tag, 0) + 1
        if tag == 'del':
            h.append(f'<p><span class="tag tag-del">删除</span><span class="del">{html_mod.escape(text)}</span></p>')
        elif tag == 'add':
            h.append(f'<p><span class="tag tag-add">新增</span><span class="add">{html_mod.escape(text)}</span></p>')
        else:
            h.append(f'<p class="same"><span class="tag" style="background:#EAECEE;color:#7F8C8D">未变</span>{html_mod.escape(text)}</p>')
h.append(f'<hr><p>统计：删除 {stats["del"]} 段 ｜ 新增 {stats["add"]} 段 ｜ 未变 {stats["same"]} 段</p>')
h.append('</body></html>')
open(OUT_HTML, 'w', encoding='utf-8').write('\n'.join(h))
print(f'HTML: {OUT_HTML}')

# ── DOCX 输出 ──
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(10.5)
doc.add_heading('V1 → V2 修订对比（红=删除，绿=新增）', level=0)
p = doc.add_paragraph('V1: docs/paper/eswa_paper.md ｜ V2: docs/paper_v2/eswa_paper_v2.md ｜ 生成日期 2026-09-14')
p.runs[0].font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

for title, entries in diff_sections:
    doc.add_heading(title, level=1)
    for tag, text in entries:
        para = doc.add_paragraph()
        r_tag = para.add_run('[删除] ' if tag == 'del' else ('[新增] ' if tag == 'add' else '[未变] '))
        r_tag.font.size = Pt(8)
        if tag == 'del':
            r_tag.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
            r_txt = para.add_run(text)
            r_txt.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
            r_txt.font.strike = True
        elif tag == 'add':
            r_tag.font.color.rgb = RGBColor(0x1E, 0x84, 0x49)
            r_txt = para.add_run(text)
            r_txt.font.color.rgb = RGBColor(0x1E, 0x84, 0x49)
        else:
            r_tag.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
            r_txt = para.add_run(text)
            r_txt.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_heading('统计', level=1)
doc.add_paragraph(f'删除 {stats["del"]} 段 ｜ 新增 {stats["add"]} 段 ｜ 未变 {stats["same"]} 段')
doc.save(OUT_DOCX)
print(f'DOCX: {OUT_DOCX}')
